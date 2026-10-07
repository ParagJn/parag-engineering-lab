"""Fill the IBM SoW Word template with the generated (and user-edited) SoW Markdown.

Only the `<< >>` marker sections and a few simple placeholders are rewritten; every
other part of the template is left exactly as authored. New paragraphs and table
rows are cloned from the template's own example content so fonts, list numbering,
and table styling carry over.
"""

import copy
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE_PATH = Path(__file__).parent / "data" / "IBM-SoW-Template.docx"

MARKER_RE = re.compile(r"<<.*?>>", re.S)
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
TABLE_WIDTH = 9420

RACI_LEGEND = "R = Responsible, A = Accountable, C = Consulted, I = Informed"
DELIVERABLES_INTRO = "The Deliverables and Work Products under this SOW are set out below."


def w(tag: str) -> str:
    return qn(f"w:{tag}")


# --- Markdown parsing --------------------------------------------------------

SECTION_ALIASES: List[Tuple[str, str]] = [
    ("out of scope", "out_of_scope"),
    ("background", "background"),
    ("deliverable", "deliverables"),
    ("work product", "deliverables"),
    ("assumption", "assumptions"),
    ("dependenc", "dependencies"),
    ("raci", "raci"),
    ("risk", "risks"),
    ("scope", "scope"),
]

Block = Tuple[str, Any]


def _section_key(heading: str) -> Optional[str]:
    h = heading.lower()
    for needle, key in SECTION_ALIASES:
        if needle in h:
            return key
    return None


def split_sections(markdown: str) -> Dict[str, List[str]]:
    """Group Markdown lines under their top-level (# / ##) section headings."""
    sections: Dict[str, List[str]] = {}
    current: Optional[List[str]] = None
    for line in markdown.splitlines():
        m = re.match(r"^#{1,2}\s+(.*)", line.strip())
        if m:
            key = _section_key(m.group(1))
            current = sections.setdefault(key, []) if key else None
            continue
        if current is not None:
            current.append(line)
    return sections


def _strip_heading_number(text: str) -> str:
    return re.sub(r"^\d+(\.\d+)*\.?\s+", "", text).strip()


def parse_blocks(lines: List[str]) -> List[Block]:
    """Turn section lines into blocks: para, bullet, h3, label, table."""
    blocks: List[Block] = []
    para: List[str] = []
    table: List[List[str]] = []

    def flush_para():
        if para:
            blocks.append(("para", " ".join(para)))
            para.clear()

    def flush_table():
        if table:
            blocks.append(("table", list(table)))
            table.clear()

    for raw in lines:
        line = raw.strip()
        if line.startswith("|"):
            flush_para()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            table.append(cells)
            continue
        flush_table()
        if not line:
            flush_para()
            continue
        m = re.match(r"^(#{3,6})\s+(.*)", line)
        if m:
            flush_para()
            kind = "h3" if len(m.group(1)) == 3 else "label"
            blocks.append((kind, _strip_heading_number(m.group(2))))
            continue
        m = re.match(r"^(?:[-*•]|\d+[.)])\s+(.*)", line)
        if m:
            flush_para()
            blocks.append(("bullet", m.group(1).strip()))
            continue
        m = re.fullmatch(r"\*\*(.+?)\*\*:?", line)
        if m:
            flush_para()
            blocks.append(("label", m.group(1).rstrip(":").strip()))
            continue
        para.append(line)
    flush_para()
    flush_table()
    return blocks


def _first_table(blocks: List[Block]) -> Tuple[List[str], List[List[str]]]:
    for kind, value in blocks:
        if kind == "table" and value:
            return value[0], value[1:]
    return [], []


def _plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).replace("`", "").strip()


def _list_items(blocks: List[Block]) -> List[str]:
    """Bullets and paragraphs as a flat list of item texts."""
    items: List[str] = []
    for kind, value in blocks:
        if kind in ("bullet", "para"):
            items.append(value)
        elif kind == "table":
            items.extend(" – ".join(c for c in row if c) for row in value[1:])
    return items


# --- Low-level XML helpers ---------------------------------------------------

def _ptext(el) -> str:
    return "".join(t.text or "" for t in el.iter(w("t")))


def _strip_ids(el) -> None:
    for node in el.iter():
        for attr in (W14 + "paraId", W14 + "textId"):
            if attr in node.attrib:
                del node.attrib[attr]


def _remove_children(parent, tag: str) -> None:
    for child in parent.findall(tag):
        parent.remove(child)


def _base_rpr(proto_p, bold: Optional[bool] = None):
    """Run properties to use for new text in a paragraph cloned from proto_p."""
    rpr = None
    for r in proto_p.iter(w("r")):
        if r.find(w("t")) is not None and r.find(w("rPr")) is not None:
            rpr = copy.deepcopy(r.find(w("rPr")))
            break
    if rpr is None:
        ppr = proto_p.find(w("pPr"))
        mark = ppr.find(w("rPr")) if ppr is not None else None
        rpr = copy.deepcopy(mark) if mark is not None else OxmlElement("w:rPr")
    for tag in ("highlight", "rStyle"):
        _remove_children(rpr, w(tag))
    if bold is not None:
        _set_bold(rpr, bold)
    return rpr


def _set_bold(rpr, bold: bool) -> None:
    _remove_children(rpr, w("b"))
    _remove_children(rpr, w("bCs"))
    if bold:
        idx = 0
        for i, child in enumerate(rpr):
            if child.tag in (w("rStyle"), w("rFonts")):
                idx = i + 1
        rpr.insert(idx, OxmlElement("w:bCs"))
        rpr.insert(idx, OxmlElement("w:b"))


def _make_run(text: str, rpr, bold: bool = False):
    r = OxmlElement("w:r")
    rp = copy.deepcopy(rpr)
    if bold:
        _set_bold(rp, True)
    r.append(rp)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def _append_rich_text(p, text: str, rpr) -> None:
    """Append text to p, honouring **bold** spans and dropping other Markdown marks."""
    text = text.replace("`", "")
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            p.append(_make_run(part.replace("*", ""), rpr, bold=(i % 2 == 1)))


def _clone_para(proto_p, text: str = "", bold: Optional[bool] = None):
    p = copy.deepcopy(proto_p)
    rpr = _base_rpr(proto_p, bold)
    for child in list(p):
        if child.tag != w("pPr"):
            p.remove(child)
    ppr = p.find(w("pPr"))
    if ppr is not None and ppr.find(w("rPr")) is not None:
        _remove_children(ppr.find(w("rPr")), w("highlight"))
    _strip_ids(p)
    if text:
        _append_rich_text(p, text, rpr)
    return p


_PPR_BEFORE_NUMPR = {"pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl"}


def _with_numbering(proto_p, num_pr):
    """Copy of proto_p carrying the given <w:numPr> (used for bullets in plain cells)."""
    p = copy.deepcopy(proto_p)
    ppr = p.find(w("pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    _remove_children(ppr, w("numPr"))
    idx = 0
    for i, child in enumerate(ppr):
        if child.tag.split("}")[1] in _PPR_BEFORE_NUMPR:
            idx = i + 1
    ppr.insert(idx, copy.deepcopy(num_pr))
    return p


def replace_in_para(p, pattern: "re.Pattern", repl: str) -> bool:
    """Replace the first match of pattern in p's text, even when it spans runs."""
    ts = list(p.iter(w("t")))
    full = "".join(t.text or "" for t in ts)
    m = pattern.search(full)
    if not m:
        return False
    start, end = m.span()
    pos = 0
    placed = False
    for t in ts:
        txt = t.text or ""
        a, b = pos, pos + len(txt)
        pos = b
        if b <= start or a >= end:
            continue
        pre = txt[: start - a] if a < start else ""
        post = txt[end - a:] if b > end else ""
        if not placed:
            t.text = pre + repl + post
            placed = True
            run_rpr = t.getparent().find(w("rPr"))
            if run_rpr is not None:
                _remove_children(run_rpr, w("highlight"))
        else:
            t.text = post
        t.set(qn("xml:space"), "preserve")
    return True


def _insert_after(anchor, elements: List[Any]) -> Any:
    for el in elements:
        anchor.addnext(el)
        anchor = el
    return anchor


def _remove(el) -> None:
    el.getparent().remove(el)


def _next_table(el):
    sib = el.getnext()
    while sib is not None and sib.tag != w("tbl"):
        sib = sib.getnext()
    return sib


# --- Table helpers -----------------------------------------------------------

def _set_cell_paras(tc, texts: List[str], bold: Optional[bool] = None) -> None:
    paras = tc.findall(w("p"))
    proto = paras[0]
    for p in paras:
        tc.remove(p)
    for text in texts or [""]:
        tc.append(_clone_para(proto, text, bold))


def _fill_table(tbl, rows: List[List[str]]) -> None:
    """Keep the header row, replace example rows with `rows` cloned from the first body row."""
    trs = tbl.findall(w("tr"))
    proto = copy.deepcopy(trs[1] if len(trs) > 1 else trs[0])
    for tr in trs[1:]:
        tbl.remove(tr)
    ncols = len(proto.findall(w("tc")))
    for row in rows:
        tr = copy.deepcopy(proto)
        _strip_ids(tr)
        for i, tc in enumerate(tr.findall(w("tc"))):
            value = row[i] if i < len(row) else ""
            if i == ncols - 1 and len(row) > ncols:
                value = " ".join(c for c in row[i:] if c)
            _set_cell_paras(tc, [value])
        tbl.append(tr)


def _build_table(proto_tbl, headers: List[str], rows: List[List[str]], widths: List[int],
                 body_cell_index: int = 0):
    """New table with the proto's table/row/cell styling but any number of columns."""
    tbl = copy.deepcopy(proto_tbl)
    _strip_ids(tbl)
    trs = tbl.findall(w("tr"))
    head_tr, body_tr = trs[0], trs[1] if len(trs) > 1 else trs[0]
    head_tc = head_tr.findall(w("tc"))[0]
    body_tcs = body_tr.findall(w("tc"))
    body_tc = body_tcs[min(body_cell_index, len(body_tcs) - 1)]

    grid = tbl.find(w("tblGrid"))
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(w("w"), str(width))
        grid.append(col)

    def make_row(src_tr, src_tc, values):
        tr = copy.deepcopy(src_tr)
        for tc in tr.findall(w("tc")):
            tr.remove(tc)
        for value, width in zip(values, widths):
            tc = copy.deepcopy(src_tc)
            tc_w = tc.find(w("tcPr")).find(w("tcW"))
            tc_w.set(w("w"), str(width))
            tc_w.set(w("type"), "dxa")
            _set_cell_paras(tc, [value])
            tr.append(tc)
        return tr

    new_head = make_row(head_tr, head_tc, headers)
    new_rows = [make_row(body_tr, body_tc, (row + [""] * len(headers))[: len(headers)]) for row in rows]
    for tr in trs:
        tbl.remove(tr)
    tbl.append(new_head)
    for tr in new_rows:
        tbl.append(tr)
    return tbl


# --- Template section fillers ------------------------------------------------

def _find_markers(body) -> Dict[str, Any]:
    markers: Dict[str, Any] = {}
    for p in body.iter(w("p")):
        text = _ptext(p).lower()
        if "<<" not in text:
            continue
        if "background" in text:
            markers["background"] = p
        elif "work product" in text or "works product" in text:
            markers["deliverables"] = p
        elif "out of scope" in text:
            markers["out_of_scope"] = p
        elif "scope" in text:
            markers["scope"] = p
        elif "assumption" in text:
            markers["assumptions"] = p
        elif "dependenc" in text:
            markers["dependencies"] = p
        elif "raci" in text:
            markers["raci"] = p
    return markers


def _fill_background(marker, blocks: List[Block], bullet_num_pr) -> None:
    elements = []
    for kind, value in blocks:
        if kind == "bullet" and bullet_num_pr is not None:
            elements.append(_clone_para(_with_numbering(marker, bullet_num_pr), value))
        elif kind == "label":
            elements.append(_clone_para(marker, value, bold=True))
        elif kind in ("para", "h3", "bullet"):
            elements.append(_clone_para(marker, value))
    if not elements:
        elements = [_clone_para(marker, "TBC")]
    _insert_after(marker, elements)
    _remove(marker)


def _fill_deliverables(marker, blocks: List[Block]) -> None:
    tbl = _next_table(marker)
    replace_in_para(marker, MARKER_RE, DELIVERABLES_INTRO)
    if tbl is None:
        return
    _, rows = _first_table(blocks)
    _fill_table(tbl, [[_plain(c) for c in row] for row in rows] or [["TBC"]])


def _fill_scope(marker, oos_marker, blocks: List[Block]) -> None:
    """Rewrite the Scope overview and replace the example subsections."""
    oos_heading = oos_marker.getprevious()
    region = []
    el = marker.getnext()
    while el is not None and el is not oos_heading:
        region.append(el)
        el = el.getnext()

    paras = [e for e in region if e.tag == w("p")]
    blank = next(p for p in paras if not _ptext(p).strip())
    numbered = [p for p in paras if p.find(".//" + w("numPr")) is not None]
    sub_heading = next(p for p in numbered if p.find(".//" + w("ilvl")).get(w("val")) == "2")
    heading_numid = sub_heading.find(".//" + w("numId")).get(w("val"))
    bullet = next(p for p in numbered if p.find(".//" + w("numId")).get(w("val")) != heading_numid)
    label = next(p for p in paras if _ptext(p).strip() in ("Scope of Work", "Overview", "Approach", "Key Outcomes"))
    body_text = next(p for p in paras if p not in numbered and p is not label
                     and _ptext(p).strip() and p.find(".//" + w("b")) is None)
    protos = {k: copy.deepcopy(v) for k, v in
              {"blank": blank, "heading": sub_heading, "bullet": bullet, "label": label, "body": body_text}.items()}

    for e in region:
        _remove(e)

    overview: List[Block] = []
    subsections: List[Tuple[str, List[Block]]] = []
    for kind, value in blocks:
        if kind == "h3":
            subsections.append((value, []))
        elif subsections:
            subsections[-1][1].append((kind, value))
        else:
            overview.append((kind, value))

    elements = []
    overview_paras = [v for k, v in overview if k == "para"]
    first_overview = overview_paras[0] if overview_paras else ("TBC" if not overview else "")
    replace_in_para(marker, MARKER_RE, _plain(first_overview))
    for kind, value in overview:
        if kind == "para" and value is first_overview:
            continue
        elements.append(_scope_element(protos, kind, value))
    elements.append(_clone_para(protos["blank"]))

    for title, sub_blocks in subsections:
        elements.append(_clone_para(protos["heading"], _plain(title)))
        for i, (kind, value) in enumerate(sub_blocks):
            if kind == "label" and i > 0:
                elements.append(_clone_para(protos["blank"]))
            elements.append(_scope_element(protos, kind, value))
        elements.append(_clone_para(protos["blank"]))

    _insert_after(marker, [e for e in elements if e is not None])


def _scope_element(protos, kind: str, value: Any):
    if kind == "label":
        return _clone_para(protos["label"], _plain(value))
    if kind == "bullet":
        return _clone_para(protos["bullet"], value)
    if kind == "para":
        return _clone_para(protos["body"], value)
    if kind == "table":
        return None
    return _clone_para(protos["body"], str(value))


def _fill_out_of_scope(marker, blocks: List[Block]) -> None:
    tbl = _next_table(marker)
    _remove(marker)
    if tbl is None:
        return
    _, rows = _first_table(blocks)
    grouped: Dict[str, List[str]] = {}
    for row in rows:
        if not row:
            continue
        area = _plain(row[0]) or "General"
        detail = " ".join(c for c in row[1:] if c).strip()
        grouped.setdefault(area, [])
        if detail:
            grouped[area].append(detail)
    if not grouped:
        grouped = {"TBC": ["Exclusions to be confirmed with the client."]}

    trs = tbl.findall(w("tr"))
    proto = copy.deepcopy(trs[1])
    for tr in trs[1:]:
        tbl.remove(tr)
    for area, details in grouped.items():
        tr = copy.deepcopy(proto)
        _strip_ids(tr)
        area_tc, detail_tc = tr.findall(w("tc"))[:2]
        _set_cell_paras(area_tc, [area])
        _set_cell_paras(detail_tc, details)
        tbl.append(tr)


def _fill_assumptions(marker, blocks: List[Block]) -> None:
    items = _list_items(blocks) or ["TBC"]
    _insert_after(marker, [_clone_para(marker, item, bold=False) for item in items])
    _remove(marker)


def _fill_dependencies(marker, blocks: List[Block]) -> None:
    tbl = _next_table(marker)
    replace_in_para(marker, MARKER_RE, "")
    if tbl is None:
        return
    _, rows = _first_table(blocks)
    _fill_table(tbl, [[_plain(c) for c in row] for row in rows] or [["TBC"]])


def _fill_raci(marker, blocks: List[Block]) -> None:
    tbl = _next_table(marker)
    replace_in_para(marker, MARKER_RE, RACI_LEGEND)
    if tbl is None:
        return
    header, rows = _first_table(blocks)
    parties = [_plain(h) for h in header[1:]] or ["IBM", "Client"]
    activity_header = _plain(header[0]) if header else "Activity"
    no_w, party_w = 600, 1100
    widths = [no_w, TABLE_WIDTH - no_w - party_w * len(parties)] + [party_w] * len(parties)
    body_rows = [[str(i)] + [_plain(c) for c in row] for i, row in enumerate(rows, 1)] or [["", "TBC"]]
    new_tbl = _build_table(tbl, ["No.", activity_header] + parties, body_rows, widths, body_cell_index=2)
    tbl.addnext(new_tbl)
    _remove(tbl)


# --- Simple placeholders -----------------------------------------------------

def parse_date(value: Any) -> Optional[date]:
    if not value:
        return None
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _fmt_date(d: Optional[date]) -> str:
    return f"{d.day} {d:%B %Y}" if d else ""


def project_date_range(tasks: Optional[List[Dict[str, Any]]]) -> Tuple[Optional[date], Optional[date]]:
    starts = [d for d in (parse_date(t.get("calculatedStartDate")) for t in tasks or []) if d]
    finishes = [d for d in (parse_date(t.get("calculatedFinishDate")) for t in tasks or []) if d]
    return (min(starts) if starts else None, max(finishes) if finishes else None)


def _fill_placeholders(body, project_name: str, start: Optional[date], end: Optional[date]) -> None:
    for el in body:
        if el.tag == w("tbl"):
            break
        if el.tag == w("p") and re.fullmatch(r"\s*\[.+\]\s*", _ptext(el)):
            replace_in_para(el, re.compile(r"\[.+\]"), project_name)
            break

    insert_date = re.compile(r"\[Insert Date: DD Month YYYY\]")
    expiry = re.compile(r"\[Insert expiry date of agreement\]")
    date_values = [start, end]
    for p in body.iter(w("p")):
        while date_values and insert_date.search(_ptext(p)):
            value = date_values.pop(0)
            if not value:
                break
            replace_in_para(p, insert_date, _fmt_date(value))
        if end and expiry.search(_ptext(p)):
            replace_in_para(p, expiry, _fmt_date(end))


def _fill_schedule(body, table_proto, tasks: Optional[List[Dict[str, Any]]]) -> None:
    if not tasks:
        return
    target = next((p for p in body.iter(w("p")) if "Insert Schedule" in _ptext(p)), None)
    if target is None:
        return
    rows = []
    for i, t in enumerate(tasks, 1):
        rows.append([
            str(t.get("index", i)),
            str(t.get("activity", "")),
            _fmt_date(parse_date(t.get("calculatedStartDate"))) or "TBD",
            _fmt_date(parse_date(t.get("calculatedFinishDate"))) or "TBD",
            str(t.get("estimatedDays", "")),
        ])
    widths = [600, 4920, 1500, 1500, 900]
    tbl = _build_table(table_proto, ["No.", "Activity", "Start Date", "Finish Date", "Days"], rows,
                       widths, body_cell_index=2)
    target.addnext(tbl)
    _remove(target)


# --- Entry point -------------------------------------------------------------

def build_sow_docx(sow_markdown: str, project_name: str,
                   tasks: Optional[List[Dict[str, Any]]] = None,
                   template_path: Path = TEMPLATE_PATH) -> bytes:
    document = docx.Document(str(template_path))
    body = document.element.body
    sections = {k: parse_blocks(v) for k, v in split_sections(sow_markdown).items()}
    markers = _find_markers(body)

    raci_proto = copy.deepcopy(_next_table(markers["raci"])) if "raci" in markers else None

    scope_marker = markers.get("scope")
    bullet_num_pr = None
    if scope_marker is not None:
        el = scope_marker.getnext()
        while el is not None and bullet_num_pr is None:
            num_pr = el.find(".//" + w("numPr"))
            if num_pr is not None and num_pr.find(w("ilvl")).get(w("val")) == "0":
                bullet_num_pr = copy.deepcopy(num_pr)
            el = el.getnext()

    if "background" in markers:
        _fill_background(markers["background"], sections.get("background", []), bullet_num_pr)
    if "deliverables" in markers:
        _fill_deliverables(markers["deliverables"], sections.get("deliverables", []))
    if scope_marker is not None and "out_of_scope" in markers:
        _fill_scope(scope_marker, markers["out_of_scope"], sections.get("scope", []))
    if "out_of_scope" in markers:
        _fill_out_of_scope(markers["out_of_scope"], sections.get("out_of_scope", []))
    if "assumptions" in markers:
        _fill_assumptions(markers["assumptions"], sections.get("assumptions", []))
    if "dependencies" in markers:
        _fill_dependencies(markers["dependencies"], sections.get("dependencies", []))
    if "raci" in markers:
        _fill_raci(markers["raci"], sections.get("raci", []))

    start, end = project_date_range(tasks)
    _fill_placeholders(body, project_name, start, end)
    if raci_proto is not None:
        _fill_schedule(body, raci_proto, tasks)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
